#!/usr/bin/env python3
"""Rebuild Bar Counter from Excel Closing Balance + replay POS sales.

1. Parse BAR COUNTER opening workbook — use **Closing Balance** as Sep 1 opening
   (Aug register end). Column D (1st AUG opening) is ignored.
2. Set Bar Counter on-hand to those openings (Product Master unchanged).
3. Clear ``stock_deducted_at`` + ``pos_invoice`` sale movements for closed
   bar/restaurant invoices from ``--from-date`` through today, then re-run
   ``deduct_stock_for_pos_invoice`` so counter = opening − sales.

Usage (repo root):
  PYTHONPATH=. .venv/bin/python scripts/rebuild_counter_from_opening_excel.py --dry-run
  PYTHONPATH=. .venv/bin/python scripts/rebuild_counter_from_opening_excel.py
  PYTHONPATH=. .venv/bin/python scripts/rebuild_counter_from_opening_excel.py \\
      --xlsx "/path/to/file.xlsx" --db /path/to/bell_elite.db
"""

from __future__ import annotations

import argparse
import os
import re
import sys
from datetime import date, datetime
from typing import Any

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if ROOT not in sys.path:
    sys.path.insert(0, ROOT)

import db as db_mod
import stores as stores_mod

DEFAULT_XLSX = (
    "/Users/rajesh/Documents/Hotel Bell elite/"
    "BAR COUNTER  OPENING STOCK FOR THE MONTH SEPTEMBER (1).xlsx"
)
DEFAULT_STOCK_REPORT_XLSX = "/Users/rajesh/Downloads/BAR STOCK REPORT (2).xlsx"
CLOSING_BALANCE_COL = 66  # 0-based: "Clossing Balance"
ITEM_NAME_COL = 2
ITEM_TYPE_COL = 1

# Excel name → Product Master name (when spellings differ).
EXCEL_NAME_ALIASES = {
    "sula (150 ml)": "Sula Satori",
    "sula 150 ml": "Sula Satori",
    "four season (150 ml)": "Four Season Red Wine",
    "four sesson white wine": "Four Sesson White Wine",
    "four sesson white white wine": "Four Sesson White Wine",
    "becardi breezer cranberry": "Becardi Breezer Cranberry",
    "becardi breezer orange": "Becardi Breezer Orange",
    "bacardi rum cranberry breezer": "Becardi Breezer Cranberry",
    "bacardi rum orange breezer": "Becardi Breezer Orange",
    "cranberry breezer": "Breezer Cranberry",
    "jamaican passion breezer": "Breezer Jamaican",
    "blackberry breezer": "Breezer Blackberry",
    "kingfisher  premium": "Kingfisher Premium",
    "absolute madrin": "Absolute Madrin",
    "absolute plain": "Absolute Plain",
    "morpheus xo": "Morpheus Xo",
    "teachers gold 12y": "Teachers Gold 12y",
    "teachers gold 12yr.": "Teachers Gold 12y",
    "teachers gold 12yr": "Teachers Gold 12y",
    "teachers highland": "Teachers Highland",
    "black and white": "Black And Whité",
    "black and whité": "Black And Whité",
    "budwiser magnum": "Budweiser Magnum",
    "budweiser magnum": "Budweiser Magnum",
    "ballantine": "Balentines",
    "mansion house": "Manson House",
    "royal palace vsop brandy": "Royal Palace Brandy",
    "royal salute": "Royal Salute 21 Years",
    "jonny walker double black": "J W Double Black",
    "jonny walker red label": "Jonny Walker Red Lable",
    "jonny walker black label": "Jonny Walker Black Lable",
    "jemesons irish whisky": "Jemison's",
    "bomore": "Bomore 12 Years",
    "warehouse tequila": "Warehouse",
    "dos flamos tequila": "Dos Flamos",
    "smirnoof plain": "Smirnoff Plain",
    "smirnoff plain": "Smirnoff Plain",
    "gray goose vodka": "Gray Goose",
    "antiquity": "Antiquity Blue",
    "jack daniels no 7": "Jack Daniels 7",
    "jack daniels fire": "Jack Daniels Fire",
}


def _normalize_match_name(name: str) -> str:
    return re.sub(r"\s+", " ", str(name or "").strip()).lower()


def _as_float(value) -> float | None:
    if value is None or value == "":
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace(",", "")
    if not text:
        return 0.0
    try:
        return float(text)
    except ValueError:
        return None


def parse_bar_counter_opening_xlsx(path: str) -> list[dict[str, Any]]:
    """Return opening lines from Closing Balance column + section unit hints."""
    from openpyxl import load_workbook

    wb = load_workbook(path, data_only=True)
    ws = wb.active
    section = ""
    unit_hint = "Bottle"
    rows: list[dict[str, Any]] = []
    for idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
        cells = list(row) if row else []
        a = cells[0] if len(cells) > 0 else None
        b = cells[ITEM_TYPE_COL] if len(cells) > ITEM_TYPE_COL else None
        c = cells[ITEM_NAME_COL] if len(cells) > ITEM_NAME_COL else None
        closing = cells[CLOSING_BALANCE_COL] if len(cells) > CLOSING_BALANCE_COL else None

        if isinstance(a, str) and a.strip() and c is None:
            section = a.strip().upper()
            if "BOTTEL" in section or "BOTTLE" in section:
                unit_hint = "Bottle"
            elif " ML" in f" {section}" or section.endswith("ML") or "IN ML" in section:
                unit_hint = "mL"
            continue

        if not isinstance(c, str) or not c.strip():
            continue
        name = c.strip()
        if name.lower() in ("item name", "clossing balance", "closing balance"):
            continue
        qty = _as_float(closing)
        if qty is None:
            continue
        rows.append(
            {
                "excel_row": idx,
                "section": section,
                "item_type": (str(b).strip() if b is not None else ""),
                "excel_name": name,
                "opening_qty": round(float(qty), 3),
                "unit_hint": unit_hint,
            }
        )
    return rows


def parse_bar_counter_additional_xlsx(path: str) -> list[dict[str, Any]]:
    """Parse Aug counter sheet: column C = product name, column BO = additional counter qty.

    Values are already in Product Master stock units (bottles / cans / mL) — add them
    on top of current Bar counter without touching sales or warehouse.
    """
    from openpyxl import load_workbook

    wb = load_workbook(path, data_only=True)
    ws = wb.active
    rows: list[dict[str, Any]] = []
    for idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
        cells = list(row) if row else []
        if idx == 1:
            continue
        name = cells[2] if len(cells) > 2 else None  # C
        if not isinstance(name, str) or not name.strip():
            continue
        excel_name = name.strip()
        if excel_name.lower() in ("item name",):
            continue
        add_qty = _as_float(cells[66] if len(cells) > 66 else None) or 0.0  # BO
        if add_qty == 0:
            continue
        category = cells[1] if len(cells) > 1 else None
        rows.append(
            {
                "excel_row": idx,
                "section": str(category or "").strip(),
                "item_type": str(category or "").strip(),
                "excel_name": excel_name,
                "opening_qty": round(float(add_qty), 3),
                "add_qty": round(float(add_qty), 3),
                "unit_hint": "Bottle",
            }
        )
    return rows


def add_bar_counter_from_additional_excel(
    conn,
    *,
    xlsx_path: str,
    dry_run: bool = False,
    user_id=None,
) -> dict[str, Any]:
    """Add Excel BO quantities onto Bar counter. Does not clear sales or warehouse."""
    db_mod.ensure_pos_schema(conn)
    db_mod.ensure_stores_schema(conn)
    opening_rows = parse_bar_counter_additional_xlsx(xlsx_path)
    match_info = match_opening_to_products(conn, opening_rows, outlet="bar")
    report: dict[str, Any] = {
        "ok": True,
        "dry_run": bool(dry_run),
        "mode": "add-counter",
        "xlsx_path": xlsx_path,
        "excel_rows": len(opening_rows),
        "matched": len(match_info["matched"]),
        "unmatched": match_info["unmatched"],
    }
    applied: list[dict[str, Any]] = []
    skipped_zero = 0
    for line in match_info["matched"]:
        add_qty = float(line.get("add_qty") or line.get("opening_qty") or 0)
        if abs(add_qty) < 0.0001:
            skipped_zero += 1
            continue
        before = float(
            stores_mod._stock_qty_on_hand(
                conn,
                "bar",
                stores_mod.STOCK_PLACE_COUNTER,
                line["product_name"],
                line["unit"],
            )
            or 0.0
        )
        if dry_run:
            applied.append(
                {
                    "product_name": line["product_name"],
                    "unit": line["unit"],
                    "add_qty": add_qty,
                    "before": before,
                    "after": round(before + add_qty, 3),
                    "excel_name": line["excel_name"],
                }
            )
            continue
        stores_mod._adjust_stock(
            conn,
            outlet="bar",
            place=stores_mod.STOCK_PLACE_COUNTER,
            item_name=line["product_name"],
            unit=line["unit"],
            qty_delta=add_qty,
            movement_type="adjustment",
            ref_type="counter_additional_import",
            ref_id=0,
            notes=f"Additional Bar counter from Excel col BO ({line['excel_name']})",
            user_id=user_id,
            allow_negative=True,
        )
        after = float(
            stores_mod._stock_qty_on_hand(
                conn,
                "bar",
                stores_mod.STOCK_PLACE_COUNTER,
                line["product_name"],
                line["unit"],
            )
            or 0.0
        )
        applied.append(
            {
                "product_name": line["product_name"],
                "unit": line["unit"],
                "add_qty": add_qty,
                "before": before,
                "after": after,
                "excel_name": line["excel_name"],
            }
        )
    report["skipped_zero"] = skipped_zero
    report["applied_count"] = len(applied)
    report["applied"] = applied
    return report


def parse_bar_stock_report_xlsx(path: str) -> list[dict[str, Any]]:
    """Parse BAR STOCK REPORT: D=counter bottles, E=warehouse bottles (as of Sep 1).

    Columns: A item, B category, C total on 01 Sept, D transfer/counter, E warehouse.
    """
    from openpyxl import load_workbook

    wb = load_workbook(path, data_only=True)
    ws = wb.active
    rows: list[dict[str, Any]] = []
    for idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
        cells = list(row) if row else []
        if idx == 1:
            continue
        name = cells[0] if len(cells) > 0 else None
        category = cells[1] if len(cells) > 1 else None
        if not isinstance(name, str) or not name.strip():
            continue
        excel_name = name.strip()
        if excel_name.lower() in ("item name",):
            continue
        counter_bottles = _as_float(cells[3] if len(cells) > 3 else None) or 0.0
        warehouse_bottles = _as_float(cells[4] if len(cells) > 4 else None) or 0.0
        total_bottles = _as_float(cells[2] if len(cells) > 2 else None)
        if total_bottles is None:
            total_bottles = counter_bottles + warehouse_bottles
        # Infer pack ml from name like "SULA (150 ML)"
        pack_ml_hint = None
        m = re.search(r"\(\s*(\d+)\s*ml\s*\)", excel_name, re.I)
        if m:
            pack_ml_hint = float(m.group(1))
        rows.append(
            {
                "excel_row": idx,
                "section": str(category or "").strip(),
                "item_type": str(category or "").strip(),
                "excel_name": excel_name,
                "counter_bottles": round(float(counter_bottles), 3),
                "warehouse_bottles": round(float(warehouse_bottles), 3),
                "total_bottles": round(float(total_bottles or 0), 3),
                "pack_ml_hint": pack_ml_hint,
                "unit_hint": "Bottle",
                # Compat with match_opening_to_products (uses opening_qty for display).
                "opening_qty": round(float(counter_bottles) + float(warehouse_bottles), 3),
            }
        )
    return rows


def _product_pack_ml(conn, product_id: int, *, pack_ml_hint: float | None = None) -> float:
    """Largest pack qty_in_base for spirits (usually 750); honor Excel (150 ML) hints."""
    if pack_ml_hint and pack_ml_hint > 0:
        return float(pack_ml_hint)
    row = conn.execute(
        """
        SELECT MAX(qty_in_base) AS qty
        FROM store_product_variants
        WHERE product_id = ?
          AND coalesce(is_active, 1) = 1
          AND coalesce(qty_in_base, 0) >= 100
        """,
        (int(product_id),),
    ).fetchone()
    try:
        qty = float((row["qty"] if row else 0) or 0)
    except (TypeError, ValueError, KeyError):
        qty = 0.0
    return qty if qty > 0 else 750.0


def _bottles_to_stock_qty(
    conn,
    *,
    product_id: int,
    default_unit: str,
    bottles: float,
    pack_ml_hint: float | None = None,
) -> float:
    """Convert Excel bottle/can counts into Product Master stock units."""
    bottles = float(bottles or 0)
    unit_key = str(default_unit or "").strip().lower()
    if unit_key in ("ml", "milliliter", "millilitre"):
        return round(bottles * _product_pack_ml(conn, product_id, pack_ml_hint=pack_ml_hint), 3)
    if unit_key in ("liter", "litre", "l"):
        # Rare catalog quirk (e.g. Breezer Cranberry) — treat Excel count as bottles of 275ml.
        return round(bottles * 0.275, 3)
    return round(bottles, 3)


def match_stock_report_to_products(
    conn, opening_rows: list[dict[str, Any]], *, outlet: str = "bar"
) -> dict[str, Any]:
    """Match stock-report lines and convert bottle counts into stock units."""
    base = match_opening_to_products(conn, opening_rows, outlet=outlet)
    matched: list[dict[str, Any]] = []
    for line in base["matched"]:
        pack_hint = line.get("pack_ml_hint")
        counter_qty = _bottles_to_stock_qty(
            conn,
            product_id=int(line["product_id"]),
            default_unit=line["unit"],
            bottles=float(line.get("counter_bottles") or 0),
            pack_ml_hint=pack_hint,
        )
        warehouse_qty = _bottles_to_stock_qty(
            conn,
            product_id=int(line["product_id"]),
            default_unit=line["unit"],
            bottles=float(line.get("warehouse_bottles") or 0),
            pack_ml_hint=pack_hint,
        )
        matched.append(
            {
                **line,
                "counter_qty": counter_qty,
                "warehouse_qty": warehouse_qty,
                "opening_qty": round(counter_qty + warehouse_qty, 3),
            }
        )
    return {
        "matched": matched,
        "unmatched": base["unmatched"],
        "product_count": base["product_count"],
    }


def match_opening_to_products(
    conn, opening_rows: list[dict[str, Any]], *, outlet: str = "bar"
) -> dict[str, Any]:
    """Match Excel names to active Product Master rows for ``outlet``."""
    outlet = stores_mod._normalize_outlet_key(outlet)
    products = conn.execute(
        """
        SELECT id, name, default_unit, outlet
        FROM store_products
        WHERE is_active = 1
          AND lower(trim(coalesce(outlet, ''))) IN (?, 'both', '')
        ORDER BY id ASC
        """,
        (outlet,),
    ).fetchall()
    by_name: dict[str, list] = {}
    for row in products:
        key = _normalize_match_name(row["name"])
        if not key:
            continue
        by_name.setdefault(key, []).append(row)

    matched: list[dict[str, Any]] = []
    unmatched: list[dict[str, Any]] = []
    for line in opening_rows:
        raw_name = line["excel_name"]
        alias = EXCEL_NAME_ALIASES.get(_normalize_match_name(raw_name))
        key = _normalize_match_name(alias or raw_name)
        candidates = by_name.get(key) or []
        # Prefer exact outlet match over both/blank.
        chosen = None
        for row in candidates:
            if stores_mod._normalize_outlet_key(row["outlet"] or outlet) == outlet:
                chosen = row
                break
        if chosen is None and candidates:
            chosen = candidates[0]
        if chosen is None:
            unmatched.append(dict(line))
            continue
        unit = (str(chosen["default_unit"] or "").strip() or line.get("unit_hint") or "pcs")
        matched.append(
            {
                **line,
                "product_id": int(chosen["id"]),
                "product_name": chosen["name"],
                "unit": unit,
                "match_key": key,
            }
        )
    return {"matched": matched, "unmatched": unmatched, "product_count": len(products)}


def _set_place_qty_absolute(
    conn,
    *,
    outlet: str,
    place: str,
    item_name: str,
    unit: str,
    target_qty: float,
    ref_type: str,
    notes: str,
    user_id=None,
) -> float:
    """Move place on-hand to ``target_qty`` via a single adjustment delta."""
    place = stores_mod._normalize_stock_place(place)
    current = stores_mod._stock_qty_on_hand(conn, outlet, place, item_name, unit)
    delta = round(float(target_qty) - float(current), 3)
    if abs(delta) < 0.0001:
        return 0.0
    return float(
        stores_mod._adjust_stock(
            conn,
            outlet=outlet,
            place=place,
            item_name=item_name,
            unit=unit,
            qty_delta=delta,
            movement_type="adjustment",
            ref_type=ref_type,
            ref_id=0,
            notes=notes,
            user_id=user_id,
            allow_shortfall=True,
            allow_negative=True,
        )
        or 0.0
    )


def _set_counter_qty_absolute(
    conn,
    *,
    outlet: str,
    item_name: str,
    unit: str,
    target_qty: float,
    ref_type: str,
    notes: str,
    user_id=None,
) -> float:
    """Move counter on-hand to ``target_qty`` via a single adjustment delta."""
    return _set_place_qty_absolute(
        conn,
        outlet=outlet,
        place=stores_mod.STOCK_PLACE_COUNTER,
        item_name=item_name,
        unit=unit,
        target_qty=target_qty,
        ref_type=ref_type,
        notes=notes,
        user_id=user_id,
    )


def apply_bar_counter_openings(
    conn,
    matched: list[dict[str, Any]],
    *,
    as_of: str,
    dry_run: bool = False,
    user_id=None,
) -> dict[str, Any]:
    """Zero all bar counter rows, then set matched openings."""
    outlet = "bar"
    place = stores_mod.STOCK_PLACE_COUNTER
    existing = conn.execute(
        """
        SELECT id, item_name, unit, qty_on_hand
        FROM store_stock_items
        WHERE outlet = ? AND place = ?
        ORDER BY id ASC
        """,
        (outlet, place),
    ).fetchall()
    applied: list[dict[str, Any]] = []
    if dry_run:
        for line in matched:
            applied.append(
                {
                    "product_name": line["product_name"],
                    "unit": line["unit"],
                    "opening_qty": line["opening_qty"],
                    "excel_name": line["excel_name"],
                }
            )
        return {
            "dry_run": True,
            "existing_bar_counter_rows": len(existing),
            "openings_to_apply": len(matched),
            "applied": applied,
        }

    # Zero every existing bar counter row first.
    for row in existing:
        _set_counter_qty_absolute(
            conn,
            outlet=outlet,
            item_name=row["item_name"],
            unit=row["unit"],
            target_qty=0.0,
            ref_type="opening_import",
            notes=f"Bar counter zero before opening import as of {as_of}",
            user_id=user_id,
        )

    for line in matched:
        delta = _set_counter_qty_absolute(
            conn,
            outlet=outlet,
            item_name=line["product_name"],
            unit=line["unit"],
            target_qty=float(line["opening_qty"]),
            ref_type="opening_import",
            notes=f"Bar counter opening as of {as_of} (Excel Closing Balance)",
            user_id=user_id,
        )
        applied.append(
            {
                "product_name": line["product_name"],
                "unit": line["unit"],
                "opening_qty": line["opening_qty"],
                "excel_name": line["excel_name"],
                "delta": round(delta, 3),
            }
        )
    return {
        "dry_run": False,
        "existing_bar_counter_rows": len(existing),
        "openings_applied": len(applied),
        "applied": applied,
    }


def apply_bar_stock_report_openings(
    conn,
    matched: list[dict[str, Any]],
    *,
    as_of: str,
    dry_run: bool = False,
    user_id=None,
) -> dict[str, Any]:
    """Zero bar counter + warehouse, then set openings from the Sep 1 stock report."""
    outlet = "bar"
    places = (stores_mod.STOCK_PLACE_COUNTER, stores_mod.STOCK_PLACE_WAREHOUSE)
    existing_by_place: dict[str, list] = {}
    for place in places:
        existing_by_place[place] = conn.execute(
            """
            SELECT id, item_name, unit, qty_on_hand
            FROM store_stock_items
            WHERE outlet = ? AND place = ?
            ORDER BY id ASC
            """,
            (outlet, place),
        ).fetchall()

    applied: list[dict[str, Any]] = []
    if dry_run:
        for line in matched:
            applied.append(
                {
                    "product_name": line["product_name"],
                    "unit": line["unit"],
                    "counter_qty": line.get("counter_qty"),
                    "warehouse_qty": line.get("warehouse_qty"),
                    "excel_name": line["excel_name"],
                    "counter_bottles": line.get("counter_bottles"),
                    "warehouse_bottles": line.get("warehouse_bottles"),
                }
            )
        return {
            "dry_run": True,
            "existing_rows": {
                p: len(existing_by_place[p]) for p in places
            },
            "openings_to_apply": len(matched),
            "applied": applied,
        }

    for place in places:
        for row in existing_by_place[place]:
            _set_place_qty_absolute(
                conn,
                outlet=outlet,
                place=place,
                item_name=row["item_name"],
                unit=row["unit"],
                target_qty=0.0,
                ref_type="opening_import",
                notes=f"Bar {place} zero before Sep 1 stock report import as of {as_of}",
                user_id=user_id,
            )

    for line in matched:
        c_delta = _set_place_qty_absolute(
            conn,
            outlet=outlet,
            place=stores_mod.STOCK_PLACE_COUNTER,
            item_name=line["product_name"],
            unit=line["unit"],
            target_qty=float(line.get("counter_qty") or 0),
            ref_type="opening_import",
            notes=f"Bar counter opening as of {as_of} (Excel col D)",
            user_id=user_id,
        )
        w_delta = _set_place_qty_absolute(
            conn,
            outlet=outlet,
            place=stores_mod.STOCK_PLACE_WAREHOUSE,
            item_name=line["product_name"],
            unit=line["unit"],
            target_qty=float(line.get("warehouse_qty") or 0),
            ref_type="opening_import",
            notes=f"Bar warehouse opening as of {as_of} (Excel col E)",
            user_id=user_id,
        )
        applied.append(
            {
                "product_name": line["product_name"],
                "unit": line["unit"],
                "counter_qty": line.get("counter_qty"),
                "warehouse_qty": line.get("warehouse_qty"),
                "excel_name": line["excel_name"],
                "counter_delta": round(c_delta, 3),
                "warehouse_delta": round(w_delta, 3),
            }
        )
    return {
        "dry_run": False,
        "existing_rows": {p: len(existing_by_place[p]) for p in places},
        "openings_applied": len(applied),
        "applied": applied,
    }


def _split_replay_invoice_ids(
    conn, invoice_ids: list[int]
) -> tuple[list[int], list[int]]:
    """Split closed invoices into bar vs restaurant (preserve date order)."""
    if not invoice_ids:
        return [], []
    placeholders = ",".join("?" for _ in invoice_ids)
    rows = conn.execute(
        f"""
        SELECT id, lower(trim(coalesce(outlet, ''))) AS outlet
        FROM pos_invoices
        WHERE id IN ({placeholders})
        ORDER BY order_date ASC, id ASC
        """,
        invoice_ids,
    ).fetchall()
    bar_ids: list[int] = []
    restaurant_ids: list[int] = []
    for row in rows:
        oid = int(row["id"])
        if row["outlet"] == "bar":
            bar_ids.append(oid)
        elif row["outlet"] == "restaurant":
            restaurant_ids.append(oid)
    return bar_ids, restaurant_ids


def clear_bar_sale_movements_for_invoices(conn, invoice_ids: list[int]) -> dict[str, int]:
    """Delete only Bar counter sale movements for these invoices (leave restaurant stock alone)."""
    if not invoice_ids:
        return {"movements_deleted": 0}
    deleted = 0
    chunk = 400
    for i in range(0, len(invoice_ids), chunk):
        part = invoice_ids[i : i + chunk]
        placeholders = ",".join("?" for _ in part)
        cur = conn.execute(
            f"""
            DELETE FROM store_stock_movements
            WHERE ref_type = 'pos_invoice'
              AND ref_id IN ({placeholders})
              AND lower(trim(coalesce(outlet, ''))) = 'bar'
            """,
            part,
        )
        deleted += int(cur.rowcount or 0)
    return {"movements_deleted": deleted}


def apply_bar_product_sales_from_restaurant_invoices(
    conn, invoice_ids: list[int], *, dry_run: bool = False, user_id=None
) -> dict[str, Any]:
    """Apply Bar-only Product Master sales from Restaurant bills onto Bar counter.

    Restaurant food stock / ``stock_deducted_at`` stay untouched so food is not
    double-deducted. Only ingredients whose product outlet routes to ``bar`` are
    applied (same rule as live ``deduct_stock_for_pos_invoice``).
    """
    if dry_run:
        return {"dry_run": True, "invoice_count": len(invoice_ids), "applied": 0}
    applied = 0
    skipped_existing = 0
    touched: list[dict[str, Any]] = []
    for invoice_id in invoice_ids:
        inv = conn.execute(
            """
            SELECT id, order_no, outlet
            FROM pos_invoices
            WHERE id = ?
            """,
            (invoice_id,),
        ).fetchone()
        if not inv:
            continue
        order_no = (inv["order_no"] or "").strip() or f"#{invoice_id}"
        lines = conn.execute(
            """
            SELECT menu_item_id, name, qty
            FROM pos_invoice_lines
            WHERE invoice_id = ?
            ORDER BY sort_order ASC, id ASC
            """,
            (invoice_id,),
        ).fetchall()
        menu_qty: dict[int, float] = {}
        for line in lines:
            mid = line["menu_item_id"]
            try:
                qty = float(line["qty"] or 0)
            except (TypeError, ValueError):
                qty = 0.0
            if mid is None or qty <= 0:
                continue
            menu_qty[int(mid)] = menu_qty.get(int(mid), 0.0) + qty
        if not menu_qty:
            continue
        recipes = db_mod.list_pos_menu_recipe_lines(conn, list(menu_qty.keys()))
        recipes_by_menu: dict[int, list] = {}
        for recipe in recipes:
            recipes_by_menu.setdefault(int(recipe["menu_item_id"]), []).append(recipe)
        needs: dict[tuple[str, str], float] = {}
        for mid, sold_qty in menu_qty.items():
            for recipe in recipes_by_menu.get(mid) or []:
                stock_outlet = stores_mod._stock_outlet_for_product(
                    recipe.get("product_outlet"), "restaurant"
                )
                if stock_outlet != "bar":
                    continue
                product_name = (recipe.get("product_name") or "").strip()
                product_unit = (recipe.get("product_unit") or "").strip() or "pcs"
                if not product_name:
                    continue
                per_portion = stores_mod._qty_in_product_units(
                    recipe.get("qty"), recipe.get("unit"), product_unit
                )
                if per_portion is None:
                    continue
                need = float(per_portion) * float(sold_qty)
                if need <= 0:
                    continue
                key = (product_name, product_unit)
                needs[key] = needs.get(key, 0.0) + need
        for (name, unit), need_qty in needs.items():
            existing = conn.execute(
                """
                SELECT id FROM store_stock_movements
                WHERE ref_type = 'pos_invoice' AND ref_id = ?
                  AND lower(trim(coalesce(outlet, ''))) = 'bar'
                  AND lower(trim(item_name)) = lower(?)
                  AND lower(trim(unit)) = lower(?)
                LIMIT 1
                """,
                (invoice_id, name, unit),
            ).fetchone()
            if existing:
                skipped_existing += 1
                continue
            stores_mod._adjust_stock(
                conn,
                outlet="bar",
                place=stores_mod.STOCK_PLACE_COUNTER,
                item_name=name,
                unit=unit,
                qty_delta=-abs(need_qty),
                movement_type="sale",
                ref_type="pos_invoice",
                ref_id=invoice_id,
                notes=f"POS sale {order_no} (restaurant→bar stock)",
                user_id=user_id,
                allow_negative=True,
            )
            applied += 1
            touched.append(
                {
                    "invoice_id": invoice_id,
                    "order_no": order_no,
                    "item_name": name,
                    "unit": unit,
                    "qty_delta": round(-abs(need_qty), 4),
                }
            )
    return {
        "dry_run": False,
        "invoice_count": len(invoice_ids),
        "applied": applied,
        "skipped_existing": skipped_existing,
        "touched": touched,
    }


def rebuild_bar_stock_from_sept_report(
    conn,
    *,
    xlsx_path: str,
    from_date: str = "2026-09-01",
    to_date: str | None = None,
    dry_run: bool = False,
    user_id=None,
) -> dict[str, Any]:
    """Set bar counter+warehouse from Sep 1 report, then replay sales after that date.

    Bar bills are cleared + fully re-deducted. Restaurant bills are not cleared
    (food stock stays intact); only Bar-product pours on those bills hit Bar counter.
    """
    db_mod.ensure_pos_schema(conn)
    db_mod.ensure_stores_schema(conn)
    to_date = to_date or date.today().isoformat()

    opening_rows = parse_bar_stock_report_xlsx(xlsx_path)
    match_info = match_stock_report_to_products(conn, opening_rows, outlet="bar")
    all_ids = list_closed_invoices_for_replay(conn, from_date=from_date, to_date=to_date)
    bar_ids, restaurant_ids = _split_replay_invoice_ids(conn, all_ids)

    report: dict[str, Any] = {
        "ok": True,
        "dry_run": bool(dry_run),
        "xlsx_path": xlsx_path,
        "from_date": from_date,
        "to_date": to_date,
        "excel_rows": len(opening_rows),
        "matched": len(match_info["matched"]),
        "unmatched": match_info["unmatched"],
        "invoice_ids": bar_ids,
        "invoice_count": len(bar_ids),
        "restaurant_invoice_ids": restaurant_ids,
        "restaurant_invoice_count": len(restaurant_ids),
    }

    if dry_run:
        report["opening_preview"] = apply_bar_stock_report_openings(
            conn, match_info["matched"], as_of=from_date, dry_run=True
        )
        report["replay_preview"] = replay_pos_stock_deductions(
            conn, bar_ids, dry_run=True
        )
        report["restaurant_bar_preview"] = apply_bar_product_sales_from_restaurant_invoices(
            conn, restaurant_ids, dry_run=True
        )
        return report

    report["opening"] = apply_bar_stock_report_openings(
        conn,
        match_info["matched"],
        as_of=from_date,
        dry_run=False,
        user_id=user_id,
    )
    report["clear"] = clear_sale_deductions_for_invoices(conn, bar_ids)
    # Drop any prior Bar sale rows tied to Restaurant bills before re-applying.
    report["clear_restaurant_bar_moves"] = clear_bar_sale_movements_for_invoices(
        conn, restaurant_ids
    )
    report["replay"] = replay_pos_stock_deductions(
        conn, bar_ids, dry_run=False, user_id=user_id
    )
    report["restaurant_bar_sales"] = apply_bar_product_sales_from_restaurant_invoices(
        conn, restaurant_ids, dry_run=False, user_id=user_id
    )

    for place in ("counter", "warehouse"):
        row = conn.execute(
            """
            SELECT COUNT(*) AS c, COALESCE(SUM(qty_on_hand), 0) AS qty
            FROM store_stock_items
            WHERE outlet = 'bar' AND place = ?
            """,
            (place,),
        ).fetchone()
        report[f"bar_{place}"] = {
            "rows": int(row["c"] or 0),
            "sum_qty": round(float(row["qty"] or 0), 3),
        }
    return report


def list_closed_invoices_for_replay(
    conn, *, from_date: str, to_date: str
) -> list[int]:
    """Closed bar/restaurant invoices in date range (exclude cancelled / import-skip)."""
    rows = conn.execute(
        """
        SELECT id
        FROM pos_invoices
        WHERE lower(trim(coalesce(outlet, ''))) IN ('bar', 'restaurant')
          AND order_date >= ?
          AND order_date <= ?
          AND trim(coalesce(cancelled_at, '')) = ''
          AND lower(trim(coalesce(status, ''))) = 'closed'
          AND lower(trim(coalesce(stock_deducted_at, ''))) != 'import-skip'
        ORDER BY order_date ASC, id ASC
        """,
        (from_date, to_date),
    ).fetchall()
    return [int(r["id"]) for r in rows]


def clear_sale_deductions_for_invoices(conn, invoice_ids: list[int]) -> dict[str, int]:
    """Delete pos_invoice sale movements and clear stock_deducted_at."""
    if not invoice_ids:
        return {"movements_deleted": 0, "invoices_cleared": 0}
    deleted = 0
    cleared = 0
    chunk = 400
    for i in range(0, len(invoice_ids), chunk):
        part = invoice_ids[i : i + chunk]
        placeholders = ",".join("?" for _ in part)
        cur = conn.execute(
            f"""
            DELETE FROM store_stock_movements
            WHERE ref_type = 'pos_invoice' AND ref_id IN ({placeholders})
            """,
            part,
        )
        deleted += int(cur.rowcount or 0)
        cur2 = conn.execute(
            f"""
            UPDATE pos_invoices
            SET stock_deducted_at = '', updated_at = datetime('now','localtime')
            WHERE id IN ({placeholders})
            """,
            part,
        )
        cleared += int(cur2.rowcount or 0)
    return {"movements_deleted": deleted, "invoices_cleared": cleared}


def replay_pos_stock_deductions(
    conn, invoice_ids: list[int], *, dry_run: bool = False, user_id=None
) -> dict[str, Any]:
    if dry_run:
        return {"dry_run": True, "invoice_count": len(invoice_ids), "replayed": 0}
    replayed = 0
    skipped = 0
    errors: list[dict[str, Any]] = []
    for invoice_id in invoice_ids:
        result = stores_mod.deduct_stock_for_pos_invoice(
            conn,
            invoice_id,
            user_id=user_id,
            allow_inactive=True,
            force_negative=True,
        )
        # ``skipped`` on success is a list of line-level skips; idempotent skip
        # uses boolean True + reason already_deducted / movements_exist.
        if result.get("ok") and result.get("skipped") is True:
            skipped += 1
        elif result.get("ok"):
            replayed += 1
        else:
            errors.append({"invoice_id": invoice_id, "result": result})
    return {
        "dry_run": False,
        "invoice_count": len(invoice_ids),
        "replayed": replayed,
        "skipped": skipped,
        "errors": errors,
    }


def rebuild_counter_from_opening_excel(
    conn,
    *,
    xlsx_path: str,
    from_date: str = "2026-09-01",
    to_date: str | None = None,
    dry_run: bool = False,
    user_id=None,
) -> dict[str, Any]:
    """Full rebuild: parse → match → open bar counter → clear → replay sales."""
    db_mod.ensure_pos_schema(conn)
    db_mod.ensure_stores_schema(conn)
    to_date = to_date or date.today().isoformat()

    opening_rows = parse_bar_counter_opening_xlsx(xlsx_path)
    match_info = match_opening_to_products(conn, opening_rows, outlet="bar")
    invoice_ids = list_closed_invoices_for_replay(
        conn, from_date=from_date, to_date=to_date
    )

    report: dict[str, Any] = {
        "ok": True,
        "dry_run": bool(dry_run),
        "xlsx_path": xlsx_path,
        "from_date": from_date,
        "to_date": to_date,
        "excel_rows": len(opening_rows),
        "matched": len(match_info["matched"]),
        "unmatched": match_info["unmatched"],
        "invoice_ids": invoice_ids,
        "invoice_count": len(invoice_ids),
    }

    if dry_run:
        report["opening_preview"] = apply_bar_counter_openings(
            conn, match_info["matched"], as_of=from_date, dry_run=True
        )
        report["replay_preview"] = replay_pos_stock_deductions(
            conn, invoice_ids, dry_run=True
        )
        return report

    report["opening"] = apply_bar_counter_openings(
        conn,
        match_info["matched"],
        as_of=from_date,
        dry_run=False,
        user_id=user_id,
    )
    report["clear"] = clear_sale_deductions_for_invoices(conn, invoice_ids)
    report["replay"] = replay_pos_stock_deductions(
        conn, invoice_ids, dry_run=False, user_id=user_id
    )

    bar_nonzero = conn.execute(
        """
        SELECT COUNT(*) AS c, COALESCE(SUM(qty_on_hand), 0) AS qty
        FROM store_stock_items
        WHERE outlet = 'bar' AND place = 'counter'
        """
    ).fetchone()
    rest_nonzero = conn.execute(
        """
        SELECT COUNT(*) AS c, COALESCE(SUM(qty_on_hand), 0) AS qty
        FROM store_stock_items
        WHERE outlet = 'restaurant' AND place = 'counter'
        """
    ).fetchone()
    report["bar_counter"] = {
        "rows": int(bar_nonzero["c"] or 0),
        "sum_qty": round(float(bar_nonzero["qty"] or 0), 3),
    }
    report["restaurant_counter"] = {
        "rows": int(rest_nonzero["c"] or 0),
        "sum_qty": round(float(rest_nonzero["qty"] or 0), 3),
    }
    report["product_master_count"] = int(
        conn.execute("SELECT COUNT(*) AS c FROM store_products").fetchone()["c"] or 0
    )
    return report


def _print_report(report: dict[str, Any]) -> None:
    prefix = "DRY-RUN " if report.get("dry_run") else ""
    print(
        f"{prefix}excel_rows={report.get('excel_rows')} matched={report.get('matched')} "
        f"unmatched={len(report.get('unmatched') or [])} "
        f"invoices={report.get('invoice_count')} "
        f"from={report.get('from_date')} to={report.get('to_date')}"
    )
    unmatched = report.get("unmatched") or []
    if unmatched:
        print("Unmatched Excel names:")
        for line in unmatched[:40]:
            print(
                f"  row {line.get('excel_row')}: {line.get('excel_name')} "
                f"opening={line.get('opening_qty')} hint={line.get('unit_hint')}"
            )
        if len(unmatched) > 40:
            print(f"  … {len(unmatched) - 40} more")

    if report.get("dry_run"):
        preview = (report.get("opening_preview") or {}).get("applied") or []
        print(f"Opening matches to apply: {len(preview)}")
        for line in preview[:25]:
            if "counter_qty" in line or "warehouse_qty" in line:
                print(
                    f"  {line['excel_name']} → {line['product_name']} ({line['unit']}) "
                    f"counter={line.get('counter_qty')} warehouse={line.get('warehouse_qty')}"
                )
            else:
                print(
                    f"  {line['excel_name']} → {line['product_name']} "
                    f"({line['unit']}) = {line['opening_qty']}"
                )
        if len(preview) > 25:
            print(f"  … {len(preview) - 25} more")
        return

    opening = report.get("opening") or {}
    clear = report.get("clear") or {}
    replay = report.get("replay") or {}
    print(
        f"openings_applied={opening.get('openings_applied')} "
        f"movements_deleted={clear.get('movements_deleted')} "
        f"invoices_cleared={clear.get('invoices_cleared')} "
        f"replayed={replay.get('replayed')} skipped={replay.get('skipped')}"
    )
    print(
        f"bar_counter rows={report.get('bar_counter', {}).get('rows')} "
        f"sum_qty={report.get('bar_counter', {}).get('sum_qty')} | "
        f"restaurant_counter rows={report.get('restaurant_counter', {}).get('rows')} "
        f"sum_qty={report.get('restaurant_counter', {}).get('sum_qty')} | "
        f"product_master={report.get('product_master_count')}"
    )
    for line in (opening.get("applied") or [])[:15]:
        if float(line.get("opening_qty") or 0) == 0:
            continue
        print(
            f"  open {line['product_name']} ({line['unit']}) = {line['opening_qty']}"
        )


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--mode",
        choices=("closing-balance", "stock-report", "add-counter"),
        default="closing-balance",
        help=(
            "closing-balance: legacy Sep opening sheet; "
            "stock-report: BAR STOCK REPORT (D=counter, E=warehouse); "
            "add-counter: add Excel col BO onto Bar counter (C=name) without touching sales"
        ),
    )
    parser.add_argument("--xlsx", default="", help="Path to opening Excel")
    parser.add_argument(
        "--db",
        default=db_mod.DATABASE_PATH,
        help="SQLite database path (default: app DATABASE_PATH)",
    )
    parser.add_argument(
        "--from-date",
        default="2026-09-01",
        help="First sales date to replay (YYYY-MM-DD)",
    )
    parser.add_argument(
        "--to-date",
        default="",
        help="Last sales date to replay (default: today)",
    )
    parser.add_argument(
        "--dry-run",
        action="store_true",
        help="Parse/match/report only; write nothing",
    )
    args = parser.parse_args()

    default_xlsx = (
        DEFAULT_STOCK_REPORT_XLSX
        if args.mode == "stock-report"
        else DEFAULT_XLSX
    )
    xlsx_path = os.path.abspath(os.path.expanduser(args.xlsx or default_xlsx))
    if not os.path.isfile(xlsx_path):
        print(f"Excel not found: {xlsx_path}", file=sys.stderr)
        return 1

    db_mod.DATABASE_PATH = os.path.abspath(os.path.expanduser(args.db))
    conn = db_mod.get_db()
    try:
        if args.mode == "add-counter":
            report = add_bar_counter_from_additional_excel(
                conn,
                xlsx_path=xlsx_path,
                dry_run=args.dry_run,
            )
            if not args.dry_run:
                conn.commit()
            print(
                f"{'DRY-RUN ' if args.dry_run else ''}"
                f"add-counter excel_rows={report.get('excel_rows')} "
                f"matched={report.get('matched')} unmatched={len(report.get('unmatched') or [])} "
                f"applied={report.get('applied_count')}"
            )
            for u in report.get("unmatched") or []:
                print(
                    f"  UNMATCH {u.get('excel_name')} "
                    f"add={u.get('add_qty') or u.get('opening_qty')}"
                )
            for line in report.get("applied") or []:
                print(
                    f"  {line['product_name']} ({line['unit']}) "
                    f"{line['before']} + {line['add_qty']} → {line['after']}"
                )
            if not args.dry_run:
                for name in ("Absolute Madrin", "Kingfisher Strong"):
                    row = conn.execute(
                        """
                        SELECT qty_on_hand, unit FROM store_stock_items
                        WHERE outlet = 'bar' AND place = 'counter'
                          AND lower(item_name) = lower(?)
                        """,
                        (name,),
                    ).fetchone()
                    if row:
                        print(f"check {name}: {row['qty_on_hand']} {row['unit']}")
            return 0 if report.get("ok") else 1

        try:
            datetime.strptime(args.from_date, "%Y-%m-%d")
        except ValueError:
            print(f"Invalid --from-date: {args.from_date}", file=sys.stderr)
            return 1
        to_date = (args.to_date or "").strip() or date.today().isoformat()
        try:
            datetime.strptime(to_date, "%Y-%m-%d")
        except ValueError:
            print(f"Invalid --to-date: {to_date}", file=sys.stderr)
            return 1

        if args.mode == "stock-report":
            report = rebuild_bar_stock_from_sept_report(
                conn,
                xlsx_path=xlsx_path,
                from_date=args.from_date,
                to_date=to_date,
                dry_run=args.dry_run,
            )
        else:
            report = rebuild_counter_from_opening_excel(
                conn,
                xlsx_path=xlsx_path,
                from_date=args.from_date,
                to_date=to_date,
                dry_run=args.dry_run,
            )
        if not args.dry_run:
            conn.commit()
        if not args.dry_run:
            sample = conn.execute(
                """
                SELECT place, item_name, unit, qty_on_hand
                FROM store_stock_items
                WHERE outlet = 'bar'
                  AND abs(coalesce(qty_on_hand, 0)) > 0.0001
                ORDER BY place ASC, lower(item_name)
                LIMIT 40
                """
            ).fetchall()
            report["bar_stock_sample"] = [dict(r) for r in sample]
    finally:
        conn.close()

    _print_report(report)
    if report.get("bar_stock_sample"):
        print("Bar stock non-zero sample:")
        for row in report["bar_stock_sample"]:
            print(
                f"  [{row['place']}] {row['item_name']} ({row['unit']}) = {row['qty_on_hand']}"
            )
    elif report.get("bar_counter_sample"):
        print("Bar counter non-zero sample:")
        for row in report["bar_counter_sample"]:
            print(f"  {row['item_name']} ({row['unit']}) = {row['qty_on_hand']}")
    if args.mode == "stock-report" and not args.dry_run:
        print(
            f"bar_counter sum={report.get('bar_counter', {}).get('sum_qty')} | "
            f"bar_warehouse sum={report.get('bar_warehouse', {}).get('sum_qty')}"
        )
    return 0 if report.get("ok") else 1


if __name__ == "__main__":
    raise SystemExit(main())
