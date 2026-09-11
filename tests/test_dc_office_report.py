"""DC Office report — bottles/pegs opening stock + July Excel two-page export."""

import os
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

import db as db_mod
from dc_office_report import (
    BEER_DATA_START,
    BEER_SECTION_LABEL_ROW,
    HEADER_TITLE_ROW,
    SPIRITS_DATA_END,
    SPIRITS_DATA_START,
    _TEMPLATE_PATH,
    _excl_bar_sale_rate,
    _qty_received_to_bottles,
    _spirits_sold_bottles_and_pegs,
    build_dc_office_opening_rows,
    build_dc_office_workbook,
    ml_to_bottles_and_pegs,
    month_label,
    resolve_bottle_ml,
)
from reports import REPORT_DEFINITIONS
from workspace_access import (
    get_endpoint_dashboard_module,
    get_endpoint_reports_submodule,
)


class DcOfficeMathTests(unittest.TestCase):
    def test_ml_split_750_bottle(self):
        bottles, pegs = ml_to_bottles_and_pegs(1580, bottle_ml=750, peg_ml=30)
        self.assertEqual(bottles, 2)
        self.assertEqual(pegs, 2)

    def test_ml_exact_bottle(self):
        bottles, pegs = ml_to_bottles_and_pegs(1500, bottle_ml=750, peg_ml=30)
        self.assertEqual(bottles, 2)
        self.assertEqual(pegs, 0)

    def test_resolve_bottle_ml_prefers_largest_pack(self):
        self.assertEqual(
            resolve_bottle_ml(
                [{"label": "375 mL", "qty_in_base": 375}, {"label": "750 mL", "qty_in_base": 750}]
            ),
            750.0,
        )
        self.assertEqual(resolve_bottle_ml([]), 750.0)
        self.assertEqual(
            resolve_bottle_ml([{"label": "700 mL", "qty_in_base": 700}]),
            700.0,
        )

    def test_qty_received_to_bottles(self):
        variants = [{"label": "750 mL", "qty_in_base": 750}]
        self.assertEqual(_qty_received_to_bottles(1500, "mL", variants), 2.0)
        self.assertEqual(_qty_received_to_bottles(4, "Bottle", variants), 4.0)
        self.assertEqual(_qty_received_to_bottles(2, "CAN", []), 2.0)

    def test_excl_bar_sale_rate_strips_vat(self):
        tax = {"vat": 0.1, "cgst": 0.025, "ugst": 0.025, "prices_include_tax": True}
        self.assertEqual(_excl_bar_sale_rate(275, tax=tax, liquor=True), 250.0)
        self.assertEqual(_excl_bar_sale_rate(330, tax=tax, liquor=True), 300.0)

    def test_spirits_sold_splits_like_opening_stock(self):
        variants = [{"label": "750 mL", "qty_in_base": 750}]
        # 25 pegs (750 ml) → 1 bottle, 0 pegs
        self.assertEqual(
            _spirits_sold_bottles_and_pegs(
                sold_ml=25 * 30, sold_bottles_raw=0, variants=variants
            ),
            (1, 0),
        )
        # 26 pegs → 1 bottle, 1 peg
        self.assertEqual(
            _spirits_sold_bottles_and_pegs(
                sold_ml=26 * 30, sold_bottles_raw=0, variants=variants
            ),
            (1, 1),
        )
        # 2 bottle serves + 2 pegs → same as 52 pegs → 2 bottles, 2 pegs
        self.assertEqual(
            _spirits_sold_bottles_and_pegs(
                sold_ml=60, sold_bottles_raw=2, variants=variants
            ),
            (2, 2),
        )


class DcOfficeDbTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.NamedTemporaryFile(suffix=".db", delete=False)
        self.tmp.close()
        self.db_path = self.tmp.name
        self._orig_path = db_mod.DATABASE_PATH
        db_mod.DATABASE_PATH = self.db_path
        db_mod.init_db()
        self.conn = db_mod.get_db()
        db_mod.ensure_stores_schema(self.conn)

        cat = self.conn.execute(
            "SELECT id FROM store_product_categories WHERE is_active = 1 ORDER BY id LIMIT 1"
        ).fetchone()
        if not cat:
            self.conn.execute(
                "INSERT INTO store_product_categories (name, sort_order, is_active) VALUES ('Whiskey', 10, 1)"
            )
            self.cat_id = self.conn.execute("SELECT last_insert_rowid()").fetchone()[0]
        else:
            self.cat_id = cat["id"]
            self.conn.execute(
                "UPDATE store_product_categories SET name = 'Whiskey' WHERE id = ?",
                (self.cat_id,),
            )

        self.conn.execute(
            """
            INSERT INTO store_product_categories (name, sort_order, is_active)
            VALUES ('Beer', 15, 1)
            """
        )
        self.beer_cat_id = self.conn.execute("SELECT last_insert_rowid()").fetchone()[0]

        self.conn.execute(
            """
            INSERT INTO store_products
                (category_id, name, default_unit, outlet, approximate_price, is_active, sort_order)
            VALUES (?, 'Absolute Plain', 'mL', 'bar', 1, 1, 1)
            """,
            (self.cat_id,),
        )
        self.ml_product_id = self.conn.execute("SELECT last_insert_rowid()").fetchone()[0]
        self.conn.execute(
            """
            INSERT INTO store_product_variants
                (product_id, label, qty_in_base, sort_order)
            VALUES (?, '750 mL', 750, 1)
            """,
            (self.ml_product_id,),
        )

        self.conn.execute(
            """
            INSERT INTO store_products
                (category_id, name, default_unit, outlet, approximate_price, is_active, sort_order)
            VALUES (?, 'Kingfisher Strong', 'Bottle', 'bar', 120, 1, 2)
            """,
            (self.beer_cat_id,),
        )
        self.bottle_product_id = self.conn.execute("SELECT last_insert_rowid()").fetchone()[0]

        self.conn.execute(
            """
            INSERT INTO store_products
                (category_id, name, default_unit, outlet, approximate_price, is_active, sort_order)
            VALUES (?, 'Zero Stock Gin', 'mL', 'bar', 1, 1, 3)
            """,
            (self.cat_id,),
        )
        self.zero_product_id = self.conn.execute("SELECT last_insert_rowid()").fetchone()[0]
        self.conn.execute(
            """
            INSERT INTO store_product_variants
                (product_id, label, qty_in_base, sort_order)
            VALUES (?, '750 mL', 750, 1)
            """,
            (self.zero_product_id,),
        )

        # Counter 1000 + warehouse 580 = 1580 ml → 2 bottles, 2 pegs
        self.conn.execute(
            """
            INSERT INTO store_stock_items (outlet, place, item_name, unit, qty_on_hand)
            VALUES ('bar', 'counter', 'Absolute Plain', 'mL', 1000)
            """
        )
        self.conn.execute(
            """
            INSERT INTO store_stock_items (outlet, place, item_name, unit, qty_on_hand)
            VALUES ('bar', 'warehouse', 'Absolute Plain', 'mL', 580)
            """
        )
        # Extra warehouse bottle of Absolute (750 mL) must fold into the same total.
        self.conn.execute(
            """
            INSERT INTO store_stock_items (outlet, place, item_name, unit, qty_on_hand)
            VALUES ('bar', 'warehouse', 'Absolute Plain', 'Bottle', 1)
            """
        )
        self.conn.execute(
            """
            INSERT INTO store_stock_items (outlet, place, item_name, unit, qty_on_hand)
            VALUES ('bar', 'counter', 'Kingfisher Strong', 'Bottle', 3)
            """
        )
        self.conn.execute(
            """
            INSERT INTO store_stock_items (outlet, place, item_name, unit, qty_on_hand)
            VALUES ('bar', 'warehouse', 'Kingfisher Strong', 'Bottle', 2)
            """
        )
        self.conn.execute(
            """
            INSERT INTO store_product_categories (name, sort_order, is_active)
            VALUES ('Alcopop', 20, 1)
            """
        )
        self.alcopop_cat_id = self.conn.execute("SELECT last_insert_rowid()").fetchone()[0]

        self.conn.execute(
            """
            INSERT INTO store_products
                (category_id, name, default_unit, outlet, approximate_price, is_active, sort_order)
            VALUES (?, 'Breezer Blackberry', 'mL', 'bar', 180, 1, 4)
            """,
            (self.alcopop_cat_id,),
        )
        self.alcopop_id = self.conn.execute("SELECT last_insert_rowid()").fetchone()[0]
        self.conn.execute(
            """
            INSERT INTO store_stock_items (outlet, place, item_name, unit, qty_on_hand)
            VALUES ('bar', 'counter', 'Breezer Blackberry', 'mL', 90)
            """
        )
        # Bar POS peg rate (incl VAT 10%): 275 → excl 250; bottle derived 250*25=6250.
        self.conn.execute(
            """
            INSERT INTO pos_menu_categories (name, sort_order, is_active, outlet)
            VALUES ('Whiskey Menu', 1, 1, 'bar')
            """
        )
        self.menu_cat_id = self.conn.execute("SELECT last_insert_rowid()").fetchone()[0]
        self.conn.execute(
            """
            INSERT INTO pos_menu_items
                (category_id, product_id, name, code, variant, rate, sort_order,
                 is_active, outlet, item_kind)
            VALUES (?, ?, 'Absolute Plain', '', '', 275, 1, 1, 'bar', 'liquor')
            """,
            (self.menu_cat_id, self.ml_product_id),
        )
        peg_menu_id = self.conn.execute("SELECT last_insert_rowid()").fetchone()[0]
        self.peg_menu_id = peg_menu_id
        self.conn.execute(
            """
            INSERT INTO pos_menu_recipe_lines (menu_item_id, product_id, qty, unit, sort_order)
            VALUES (?, ?, 30, 'ml', 1)
            """,
            (peg_menu_id, self.ml_product_id),
        )
        # Absolute bottle serve (1 bottle) for sale column H.
        self.conn.execute(
            """
            INSERT INTO pos_menu_items
                (category_id, product_id, name, code, variant, rate, sort_order,
                 is_active, outlet, item_kind)
            VALUES (?, ?, 'Absolute Plain Bottle', '', '', 6875, 3, 1, 'bar', 'liquor')
            """,
            (self.menu_cat_id, self.ml_product_id),
        )
        self.abs_bottle_menu_id = self.conn.execute("SELECT last_insert_rowid()").fetchone()[0]
        self.conn.execute(
            """
            INSERT INTO pos_menu_recipe_lines (menu_item_id, product_id, qty, unit, sort_order)
            VALUES (?, ?, 1, 'bottle', 1)
            """,
            (self.abs_bottle_menu_id, self.ml_product_id),
        )
        self.conn.execute(
            """
            INSERT INTO pos_menu_items
                (category_id, product_id, name, code, variant, rate, sort_order,
                 is_active, outlet, item_kind)
            VALUES (?, ?, 'Kingfisher Strong', '', '', 330, 2, 1, 'bar', 'liquor')
            """,
            (self.menu_cat_id, self.bottle_product_id),
        )
        beer_menu_id = self.conn.execute("SELECT last_insert_rowid()").fetchone()[0]
        self.beer_menu_id = beer_menu_id
        self.conn.execute(
            """
            INSERT INTO pos_menu_recipe_lines (menu_item_id, product_id, qty, unit, sort_order)
            VALUES (?, ?, 1, 'bottle', 1)
            """,
            (beer_menu_id, self.bottle_product_id),
        )
        # September sales: 8 pegs + 2 bottles Absolute; 7 Kingfisher bottles.
        self.conn.execute(
            """
            INSERT INTO pos_invoices
                (order_no, saved_at, order_date, outlet, status, is_active,
                 customer_name, subtotal, grand_total, created_at, updated_at)
            VALUES
                ('DCO-1', '2026-09-05 20:00:00', '2026-09-05', 'bar', 'closed', 0,
                 'Guest', 100, 100, '2026-09-05 20:00:00', '2026-09-05 20:00:00')
            """
        )
        inv1 = self.conn.execute("SELECT last_insert_rowid()").fetchone()[0]
        self.conn.execute(
            """
            INSERT INTO pos_invoice_lines
                (invoice_id, menu_item_id, name, variant, rate, qty, line_total, sort_order)
            VALUES
                (?, ?, 'Absolute Plain', '', 275, 8, 2200, 1),
                (?, ?, 'Absolute Plain Bottle', '', 6875, 2, 13750, 2),
                (?, ?, 'Kingfisher Strong', '', 330, 7, 2310, 3)
            """,
            (inv1, self.peg_menu_id, inv1, self.abs_bottle_menu_id, inv1, self.beer_menu_id),
        )
        # Out-of-month sale ignored.
        self.conn.execute(
            """
            INSERT INTO pos_invoices
                (order_no, saved_at, order_date, outlet, status, is_active,
                 customer_name, subtotal, grand_total, created_at, updated_at)
            VALUES
                ('DCO-AUG', '2026-08-20 20:00:00', '2026-08-20', 'bar', 'closed', 0,
                 'Guest', 275, 275, '2026-08-20 20:00:00', '2026-08-20 20:00:00')
            """
        )
        inv_aug = self.conn.execute("SELECT last_insert_rowid()").fetchone()[0]
        self.conn.execute(
            """
            INSERT INTO pos_invoice_lines
                (invoice_id, menu_item_id, name, variant, rate, qty, line_total, sort_order)
            VALUES (?, ?, 'Absolute Plain', '', 275, 99, 27225, 1)
            """,
            (inv_aug, self.peg_menu_id),
        )
        # Restaurant bill of bar liquor must still count (Absolute Madrin-style).
        self.conn.execute(
            """
            INSERT INTO pos_invoices
                (order_no, saved_at, order_date, outlet, status, is_active,
                 customer_name, subtotal, grand_total, created_at, updated_at,
                 settled_at)
            VALUES
                ('DCO-REST', '2026-09-01 22:00:00', '2026-09-01', 'restaurant', 'closed', 0,
                 'Guest', 550, 550, '2026-09-01 22:00:00', '2026-09-01 22:00:00',
                 '2026-09-01 22:58:00')
            """
        )
        inv_rest = self.conn.execute("SELECT last_insert_rowid()").fetchone()[0]
        self.conn.execute(
            """
            INSERT INTO pos_invoice_lines
                (invoice_id, menu_item_id, name, variant, rate, qty, line_total, sort_order)
            VALUES (?, ?, 'Absolute Plain', '', 275, 2, 550, 1)
            """,
            (inv_rest, self.peg_menu_id),
        )
        # 60 ml double serve → 2 pegs per sale.
        self.conn.execute(
            """
            INSERT INTO pos_menu_items
                (category_id, product_id, name, code, variant, rate, sort_order,
                 is_active, outlet, item_kind)
            VALUES (?, ?, 'Absolute Plain Double', '', '', 500, 4, 1, 'bar', 'liquor')
            """,
            (self.menu_cat_id, self.ml_product_id),
        )
        self.double_menu_id = self.conn.execute("SELECT last_insert_rowid()").fetchone()[0]
        self.conn.execute(
            """
            INSERT INTO pos_menu_recipe_lines (menu_item_id, product_id, qty, unit, sort_order)
            VALUES (?, ?, 60, 'ml', 1)
            """,
            (self.double_menu_id, self.ml_product_id),
        )
        self.conn.execute(
            """
            INSERT INTO pos_invoices
                (order_no, saved_at, order_date, outlet, status, is_active,
                 customer_name, subtotal, grand_total, created_at, updated_at)
            VALUES
                ('DCO-60', '2026-09-08 21:00:00', '2026-09-08', 'bar', 'closed', 0,
                 'Guest', 500, 500, '2026-09-08 21:00:00', '2026-09-08 21:00:00')
            """
        )
        inv60 = self.conn.execute("SELECT last_insert_rowid()").fetchone()[0]
        self.conn.execute(
            """
            INSERT INTO pos_invoice_lines
                (invoice_id, menu_item_id, name, variant, rate, qty, line_total, sort_order)
            VALUES (?, ?, 'Absolute Plain Double', '', 500, 1, 500, 1)
            """,
            (inv60, self.double_menu_id),
        )
        # September purchases: 1500 mL Absolute (= 2 bottles), 4 Kingfisher bottles.
        # Out-of-month receive must be ignored.
        self.conn.execute(
            """
            INSERT INTO store_stock_movements
                (outlet, place, item_name, unit, qty_delta, movement_type, ref_type, created_at)
            VALUES
                ('bar', 'warehouse', 'Absolute Plain', 'mL', 1500, 'receive',
                 'stock_inward_direct', '2026-09-10 10:00:00'),
                ('bar', 'warehouse', 'Kingfisher Strong', 'Bottle', 4, 'receive',
                 'stock_inward_direct', '2026-09-12 11:00:00'),
                ('bar', 'warehouse', 'Absolute Plain', 'mL', 750, 'receive',
                 'stock_inward_direct', '2026-08-15 10:00:00'),
                ('bar', 'counter', 'Absolute Plain', 'mL', 750, 'receive',
                 'stock_transfer', '2026-09-20 10:00:00')
            """
        )
        self.conn.commit()

    def tearDown(self):
        try:
            self.conn.close()
        except Exception:
            pass
        db_mod.DATABASE_PATH = self._orig_path
        try:
            os.unlink(self.db_path)
        except OSError:
            pass

    def test_spirits_ml_split_and_beer_bottles_no_ml_convert(self):
        rows = build_dc_office_opening_rows(self.conn, year=2026, month=9)
        by_name = {r["name"]: r for r in rows}
        self.assertIn("Absolute Plain", by_name)
        self.assertIn("Kingfisher Strong", by_name)
        self.assertIn("Zero Stock Gin", by_name)

        abs_row = by_name["Absolute Plain"]
        self.assertEqual(abs_row["page"], "spirits")
        # 1000 counter mL + 580 warehouse mL + 1 warehouse Bottle(750) = 2330 mL
        self.assertEqual(abs_row["bottles"], 3)
        self.assertEqual(abs_row["pegs_30ml"], 2)
        self.assertEqual(abs_row["purchased_bottles"], 2)
        # 275 incl VAT 10% → 250 excl peg; bottle = 250 * 25
        self.assertEqual(abs_row["rate_per_peg"], 250.0)
        self.assertEqual(abs_row["rate_per_bottle"], 6250.0)
        self.assertEqual(abs_row["sold_bottles"], 2)
        # Total sold mL = 8*30 + 2*30 + 60 + 2*750 = 1860 → 2 bottles, 12 pegs
        # (same split rule as Opening Stock: 750 ml = 25 pegs/bottle)
        self.assertEqual(abs_row["sold_pegs"], 12)

        kf = by_name["Kingfisher Strong"]
        self.assertEqual(kf["page"], "beer")
        # 3 counter + 2 warehouse bottles
        self.assertEqual(kf["bottles"], 5)
        self.assertEqual(kf["pegs_30ml"], 0)
        self.assertEqual(kf["purchased_bottles"], 4)
        self.assertEqual(kf["rate_per_bottle"], 300.0)  # 330 / 1.1
        self.assertEqual(kf["sold_bottles"], 7)
        self.assertEqual(kf["sold_pegs"], 0)

        zero = by_name["Zero Stock Gin"]
        self.assertEqual(zero["page"], "spirits")
        self.assertEqual(zero["bottles"], 0)
        self.assertEqual(zero["pegs_30ml"], 0)
        self.assertEqual(zero["purchased_bottles"], 0)

        # Alcopop category is bottle-sold even if unit is mL — no peg split.
        alcopop = by_name["Breezer Blackberry"]
        self.assertEqual(alcopop["page"], "beer")
        self.assertEqual(alcopop["bottles"], 90)
        self.assertEqual(alcopop["pegs_30ml"], 0)
        self.assertEqual(alcopop["purchased_bottles"], 0)

        # Outside selected month → no purchased bottles counted.
        aug = build_dc_office_opening_rows(self.conn, year=2026, month=8)
        self.assertEqual({r["name"]: r["purchased_bottles"] for r in aug}["Absolute Plain"], 1)
        self.assertEqual({r["name"]: r["purchased_bottles"] for r in aug}["Kingfisher Strong"], 0)

    def test_workbook_july_format_two_pages_clears_template_values(self):
        self.assertTrue(_TEMPLATE_PATH.is_file(), f"missing template {_TEMPLATE_PATH}")
        wb0 = load_workbook(_TEMPLATE_PATH)
        ws0 = wb0.active
        ws0.cell(SPIRITS_DATA_START, 2).value = "STALE_SPIRIT"
        ws0.cell(SPIRITS_DATA_START, 3).value = 999
        ws0.cell(SPIRITS_DATA_START, 5).value = 55
        ws0.cell(BEER_DATA_START, 2).value = "STALE_BEER"
        ws0.cell(BEER_DATA_START, 3).value = 888
        tmp_xlsx = Path(self.db_path).with_suffix(".xlsx")
        wb0.save(tmp_xlsx)

        rows = build_dc_office_opening_rows(self.conn, year=2026, month=9)
        wb = build_dc_office_workbook(
            rows,
            month_label_text=month_label(2026, 9),
            template_path=tmp_xlsx,
        )
        self.assertEqual(wb.sheetnames, ["Spirits", "Beer Alcopop"])
        ws1 = wb["Spirits"]
        ws2 = wb["Beer Alcopop"]

        title = str(ws1.cell(HEADER_TITLE_ROW, 1).value or "")
        self.assertIn("SEPTEMBER 2026", title.upper())
        self.assertEqual(ws1.cell(1, 1).value, "HOTEL BELL ELITE")
        self.assertEqual(ws1.cell(5, 1).value, "S/No")
        self.assertEqual(ws1.cell(6, 3).value, "Bottel")

        spirits = [r for r in rows if r["page"] == "spirits"]
        beer = [r for r in rows if r["page"] == "beer"]
        self.assertTrue(spirits)
        self.assertTrue(beer)

        first_spirit = spirits[0]
        self.assertEqual(ws1.cell(SPIRITS_DATA_START, 1).value, first_spirit["sno"])
        self.assertEqual(ws1.cell(SPIRITS_DATA_START, 2).value, first_spirit["name"])
        self.assertEqual(ws1.cell(SPIRITS_DATA_START, 3).value, first_spirit["bottles"])
        self.assertEqual(ws1.cell(SPIRITS_DATA_START, 4).value, first_spirit["pegs_30ml"])
        self.assertEqual(
            ws1.cell(SPIRITS_DATA_START, 5).value, first_spirit["purchased_bottles"]
        )
        self.assertEqual(
            ws1.cell(SPIRITS_DATA_START, 6).value, f"=C{SPIRITS_DATA_START}+E{SPIRITS_DATA_START}"
        )
        self.assertNotEqual(ws1.cell(SPIRITS_DATA_START, 2).value, "STALE_SPIRIT")

        # Sheet 1 must not list beer/alcopop products.
        for r in range(SPIRITS_DATA_START, SPIRITS_DATA_END + 1):
            name = ws1.cell(r, 2).value
            if not name:
                continue
            self.assertNotIn(
                str(name),
                {b["name"] for b in beer},
                f"beer product {name!r} still on sheet 1 row {r}",
            )
        self.assertIsNone(ws1.cell(BEER_SECTION_LABEL_ROW, 1).value)
        self.assertIsNone(ws1.cell(BEER_DATA_START, 2).value)

        # Sheet 2 holds Beer / Alcopop.
        self.assertEqual(ws2.cell(7, 1).value, "BEER / ALCOPOP")
        first_beer = beer[0]
        self.assertEqual(ws2.cell(8, 1).value, first_beer["sno"])
        self.assertEqual(ws2.cell(8, 2).value, first_beer["name"])
        self.assertEqual(ws2.cell(8, 3).value, first_beer["bottles"])
        self.assertEqual(ws2.cell(8, 4).value, 0)
        self.assertEqual(ws2.cell(8, 5).value, first_beer["purchased_bottles"])
        self.assertEqual(ws2.cell(8, 6).value, "=C8+E8")
        # Beer sheet drops empty J/K: rate at J, amount at K, discard L/M, closing N/O.
        self.assertEqual(ws2.cell(6, 10).value, "rate per btl")
        self.assertEqual(ws2.cell(6, 11).value, "total amt")
        self.assertIsNone(ws2.cell(6, 16).value)
        self.assertEqual(ws2.cell(8, 10).value, first_beer["rate_per_bottle"])
        self.assertEqual(ws2.cell(8, 11).value, "=H8*J8")
        self.assertEqual(ws2.cell(8, 14).value, "=F8-H8-L8")
        self.assertIsNone(ws2.cell(8, 16).value)
        self.assertNotEqual(ws2.cell(8, 2).value, "STALE_BEER")
        # Spirits: J = bottle rate, L = peg rate (excl tax); K/M formulas updated.
        abs_spirit = next(r for r in spirits if r["name"] == "Absolute Plain")
        abs_row_num = SPIRITS_DATA_START + spirits.index(abs_spirit)
        self.assertEqual(ws1.cell(6, 10).value, "Rate per bottle")
        self.assertEqual(ws1.cell(6, 12).value, "Rate per peg")
        self.assertEqual(ws1.cell(abs_row_num, 10).value, abs_spirit["rate_per_bottle"])
        self.assertEqual(ws1.cell(abs_row_num, 12).value, abs_spirit["rate_per_peg"])
        self.assertEqual(ws1.cell(abs_row_num, 8).value, abs_spirit["sold_bottles"])
        self.assertEqual(ws1.cell(abs_row_num, 9).value, abs_spirit["sold_pegs"])
        self.assertEqual(
            ws1.cell(abs_row_num, 11).value,
            f"=I{abs_row_num}+H{abs_row_num}*25",
        )
        self.assertEqual(
            ws1.cell(abs_row_num, 13).value,
            f"=H{abs_row_num}*J{abs_row_num}+I{abs_row_num}*L{abs_row_num}",
        )
        self.assertEqual(
            ws1.cell(abs_row_num, 16).value,
            f"=F{abs_row_num}-H{abs_row_num}-N{abs_row_num}",
        )
        self.assertEqual(
            ws1.cell(abs_row_num, 17).value,
            f"=G{abs_row_num}-I{abs_row_num}-O{abs_row_num}",
        )
        kf_beer = next(r for r in beer if r["name"] == "Kingfisher Strong")
        kf_row_num = 8 + beer.index(kf_beer)
        self.assertEqual(ws2.cell(kf_row_num, 8).value, kf_beer["sold_bottles"])

        try:
            tmp_xlsx.unlink()
        except OSError:
            pass


class DcOfficeHubWiringTests(unittest.TestCase):
    def test_card_and_endpoints_registered(self):
        by_id = {r["id"]: r for r in REPORT_DEFINITIONS}
        self.assertIn("dc_office", by_id)
        self.assertEqual(by_id["dc_office"]["name"], "DC Office")
        self.assertEqual(by_id["dc_office"]["view_route"], "sales_report_dc_office")
        self.assertEqual(
            by_id["dc_office"]["download_route"], "sales_report_dc_office_export"
        )
        self.assertEqual(
            get_endpoint_dashboard_module("sales_report_dc_office"), "reports"
        )
        self.assertEqual(
            get_endpoint_reports_submodule("sales_report_dc_office"), "dc_office"
        )
        self.assertEqual(
            get_endpoint_reports_submodule("sales_report_dc_office_export"),
            "dc_office",
        )


if __name__ == "__main__":
    unittest.main()
