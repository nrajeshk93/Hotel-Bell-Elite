"""Monthly Payroll export includes EPF-exempt employees on the primary sheet."""

import os
import tempfile
import unittest
from io import BytesIO
from unittest import mock

from openpyxl import load_workbook

import db as db_mod


class MonthlyPayrollExportEpfExemptTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.NamedTemporaryFile(suffix=".db", delete=False)
        self.tmp.close()
        self.db_path = self.tmp.name
        self._orig_path = db_mod.DATABASE_PATH
        db_mod.DATABASE_PATH = self.db_path
        db_mod.init_db()

        import app as app_mod
        import employee_payroll as payroll_mod

        self.app_mod = app_mod
        self.payroll_mod = payroll_mod
        self.app = app_mod.app
        self.app.config["TESTING"] = True
        self.client = self.app.test_client()

        conn = db_mod.get_db()
        try:
            admin = conn.execute(
                "SELECT id FROM users WHERE username = 'admin'"
            ).fetchone()
            self.admin_id = admin["id"]
            conn.executemany(
                """INSERT INTO employees
                   (emp_code, name, location, mobile, status, gross_salary,
                    epf_exempt, esic_exempt, account_number, ifsc_code)
                   VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""",
                [
                    (
                        "HBE1",
                        "EPF Worker",
                        "FO",
                        "9000000001",
                        "active",
                        20000,
                        0,
                        0,
                        "111",
                        "ICIC0001111",
                    ),
                    (
                        "HBE27",
                        "SUNIL RAWAT",
                        "KITCHEN",
                        "9000000027",
                        "active",
                        30000,
                        1,
                        1,
                        "1592000100081821",
                        "PUNB0159200",
                    ),
                    (
                        "HBE99",
                        "Inactive Exempt",
                        "HK",
                        "9000000099",
                        "inactive",
                        15000,
                        1,
                        1,
                        "222",
                        "SBIN0002222",
                    ),
                ],
            )
            conn.commit()
        finally:
            conn.close()

        self.user = {
            "id": self.admin_id,
            "username": "admin",
            "full_name": "Administrator",
            "is_admin": True,
            "is_active": True,
            "dashboard_access": set(),
            "payroll_access": set(),
        }
        self._patches = [
            mock.patch.object(app_mod, "get_current_user", return_value=self.user),
            mock.patch.object(payroll_mod, "get_current_user", return_value=self.user),
            mock.patch.object(
                payroll_mod, "_default_reporting_period", return_value=(2026, 8)
            ),
        ]
        for p in self._patches:
            p.start()

    def tearDown(self):
        for p in self._patches:
            p.stop()
        db_mod.DATABASE_PATH = self._orig_path
        try:
            os.unlink(self.db_path)
        except OSError:
            pass

    def test_export_primary_sheet_includes_epf_exempt_active(self):
        resp = self.client.get("/export_employees?year=2026&month=8")
        self.assertEqual(resp.status_code, 200)
        wb = load_workbook(BytesIO(resp.data))
        self.assertIn("Hotel Bell Elite", wb.sheetnames)

        primary = wb["Hotel Bell Elite"]
        primary_codes = [
            str(row[0])
            for row in primary.iter_rows(min_row=2, max_col=1, values_only=True)
            if row[0]
        ]
        self.assertIn("HBE27", primary_codes)
        self.assertIn("HBE1", primary_codes)
        self.assertNotIn("HBE99", primary_codes)

        self.assertIn("Non EPF Employees", wb.sheetnames)
        non_epf = wb["Non EPF Employees"]
        non_epf_codes = [
            str(row[0])
            for row in non_epf.iter_rows(min_row=2, max_col=1, values_only=True)
            if row[0]
        ]
        self.assertEqual(non_epf_codes, ["HBE27"])

    def test_monthly_payroll_view_includes_epf_exempt_active(self):
        page = self.client.get("/monthly_payroll?year=2026&month=8")
        self.assertEqual(page.status_code, 200)
        html = page.get_data(as_text=True)
        self.assertIn("HBE27", html)
        self.assertIn("SUNIL RAWAT", html)

    def test_export_present_matches_ui_badge_with_half_days(self):
        """Present column must use P + ½H (UI badge), not raw full-day count."""
        conn = db_mod.get_db()
        try:
            emp = conn.execute(
                "SELECT id FROM employees WHERE emp_code = 'HBE1'"
            ).fetchone()
            emp_id = emp["id"]
            # August 2026: 28 full present + 2 half days → UI badge 29, not 28.
            for day in range(1, 29):
                conn.execute(
                    "INSERT INTO attendance (employee_id, date, status, updated_at) "
                    "VALUES (?, ?, 'present', '2026-08-28')",
                    (emp_id, f"2026-08-{day:02d}"),
                )
            for day in (29, 30):
                conn.execute(
                    "INSERT INTO attendance (employee_id, date, status, updated_at) "
                    "VALUES (?, ?, 'half_day', '2026-08-30')",
                    (emp_id, f"2026-08-{day:02d}"),
                )
            conn.commit()
        finally:
            conn.close()

        page = self.client.get("/monthly_payroll?year=2026&month=8")
        self.assertEqual(page.status_code, 200)
        html = page.get_data(as_text=True)
        self.assertIn("29/31", html)

        resp = self.client.get("/export_employees?year=2026&month=8")
        self.assertEqual(resp.status_code, 200)
        wb = load_workbook(BytesIO(resp.data))
        primary = wb["Hotel Bell Elite"]
        headers = [cell.value for cell in primary[1]]
        present_col = headers.index("Present") + 1
        present_by_code = {}
        for row in primary.iter_rows(min_row=2, values_only=True):
            if row[0]:
                present_by_code[str(row[0])] = row[present_col - 1]
        self.assertEqual(present_by_code["HBE1"], 29)


if __name__ == "__main__":
    unittest.main()
